import os
from docx import Document
from openpyxl import load_workbook
import PyPDF2

def extraer_metadatos_docx(ruta_archivo):
    try:
        doc = Document(ruta_archivo)
        propiedades = doc.core_properties
        metadatos = {
            "autor": propiedades.author,
            "titulo": propiedades.title,
            "tema": propiedades.subject,
            "comentarios": propiedades.comments,
            "palabras_clave": propiedades.keywords,
        }
        return metadatos
    except Exception as e:
        print(f"Error al extraer metadatos de {ruta_archivo}: {e}")
        return {}

def extraer_metadatos_xlsx(ruta_archivo):
    try:
        wb = load_workbook(filename=ruta_archivo, data_only=True)
        propiedades = wb.properties
        metadatos = {
            "autor": propiedades.creator,
            "titulo": propiedades.title,
            "tema": propiedades.subject,
            "comentarios": propiedades.comments,
            "palabras_clave": propiedades.keywords,
        }
        return metadatos
    except Exception as e:
        print(f"Error al extraer metadatos de {ruta_archivo}: {e}")
        return {}

import PyPDF2

def extraer_metadatos_pdf(ruta_archivo):
    try:
        with open(ruta_archivo, 'rb') as archivo_pdf:
            lector_pdf = PyPDF2.PdfReader(archivo_pdf)
            info = lector_pdf.metadata
            metadatos = {
                "autor": info.get('/Author', None),
                "titulo": info.get('/Title', None),
                "tema": info.get('/Subject', None),
                "productor": info.get('/Producer', None),
                "palabras_clave": info.get('/Keywords', None),
            }
            return metadatos
    except Exception as e:
        print(f"Error al extraer metadatos de {ruta_archivo}: {e}")
        return {}

def main():
    # Rutas
    ruta_docx = "/home/kalitony/Documents/PracticaDos/2.docx"
    ruta_xlsx = "/home/kalitony/Documents/PracticaDos/archivo.xlsx"
    ruta_pdf = "/home/kalitony/Documents/PracticaDos/1.pdf"

    while True:
        print("\nSeleccione el tipo de archivo para analizar los metadatos:")
        print("1. DOCX")
        print("2. XLSX")
        print("3. PDF")
        print("4. Salir")
        opcion = input("Ingrese el número de su opción: ")

        if opcion == "1":
            print("Metadatos DOCX:", extraer_metadatos_docx(ruta_docx))
        elif opcion == "2":
            print("Metadatos XLSX:", extraer_metadatos_xlsx(ruta_xlsx))
        elif opcion == "3":
            print("Metadatos PDF:", extraer_metadatos_pdf(ruta_pdf))
        elif opcion == "4":
            print("Saliendo...")
            break
        else:
            print("Opción no válida. Por favor, intente nuevamente.")

if __name__ == "__main__":
    main()

